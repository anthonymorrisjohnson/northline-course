"""Rebuild slides/present.html as an editable PowerPoint: native text, shapes and tables.

scripts/export_deck.py makes a pixel-exact deck of pictures. This one trades the web fonts
for Arial and Courier New so that every word can be edited in PowerPoint, Keynote or
Google Slides. Author-only:

    uv run --with python-pptx --with beautifulsoup4 --with pillow python scripts/export_deck_editable.py

Text boxes are sized by measuring real line wrapping with Arial and Courier New when those
fonts are on the machine, and by a character-count estimate when they are not.
"""
from __future__ import annotations

import math
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from export_deck import slides_meta  # noqa: E402

SOURCE = ROOT / "slides" / "present.html"
OUT = ROOT / "slides" / "present-editable.pptx"

# The browser stage is 1600 x 900 px on a 13.333 x 7.5 in slide: 120 px to the inch.
EMU_PER_PX = 7620
PT_PER_PX = 0.6
W, H = 1600, 900
LEFT, TOP, RIGHT, BOTTOM, GAP = 112, 86, 112, 128, 30
CONTENT_W = W - LEFT - RIGHT

SANS, MONO = "Arial", "Courier New"
LIGHT = {"bg": "F2F5F7", "ink": "13212C", "muted": "566977", "accent": "0B6E8A", "rail": "CBD5DC"}
DARK = {"bg": "13212C", "ink": "F2F5F7", "muted": "9DB0BD", "accent": "6CC3DB", "rail": "2A3A46"}
CARD, RULE, ALERT, ALERT_SOFT, INK, PAPER = "FFFFFF", "CBD5DC", "B3261E", "F7DEDB", "13212C", "F2F5F7"


# ---------- text model: a paragraph is a list of runs, a run is (text, overrides) ----------

def runs_of(node, base: dict | None = None) -> list[list[tuple[str, dict]]]:
    """Inline HTML to paragraphs of runs. <br> starts a paragraph; b, .mono and .hot restyle a run."""
    paras: list[list[tuple[str, dict]]] = [[]]

    def walk(n, style):
        for child in n.children:
            if isinstance(child, str):
                text = re.sub(r"\s+", " ", child)
                if text:
                    paras[-1].append((text, style))
            elif child.name == "br":
                paras.append([])
            else:
                s = dict(style)
                classes = child.get("class") or []
                if child.name == "b":
                    s["bold"] = True
                if "mono" in classes:
                    s["font"] = MONO
                if "hot" in classes:
                    s.update(color=ALERT, bold=True)
                walk(child, s)

    walk(node, dict(base or {}))
    out = []
    for p in paras:
        if p:
            p[0] = (p[0][0].lstrip(), p[0][1])
            p[-1] = (p[-1][0].rstrip(), p[-1][1])
        out.append([r for r in p if r[0]])
    return [p for p in out if p] or [[("", {})]]


def plain(paras) -> str:
    return "\n".join("".join(t for t, _ in p) for p in paras)


FONT_FILES = {
    (SANS, False): ["/System/Library/Fonts/Supplemental/Arial.ttf", "C:/Windows/Fonts/arial.ttf",
                    "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf"],
    (SANS, True): ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "C:/Windows/Fonts/arialbd.ttf",
                   "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf"],
    (MONO, False): ["/System/Library/Fonts/Supplemental/Courier New.ttf", "C:/Windows/Fonts/cour.ttf",
                    "/usr/share/fonts/truetype/msttcorefonts/Courier_New.ttf"],
}
_FONTS: dict = {}


def _measure(font: str, bold: bool):
    """A function text -> width in em, from the real font when it is installed."""
    key = (font, bold and font == SANS)
    if key not in _FONTS:
        _FONTS[key] = None
        try:
            from PIL import ImageFont
            for path in FONT_FILES[key]:
                if Path(path).exists():
                    face = ImageFont.truetype(path, 200)
                    _FONTS[key] = lambda t, f=face: f.getlength(t) / 200
                    break
        except ImportError:
            pass
    return _FONTS[key]


def lines_needed(paras, size: float, width: float, font: str = SANS, bold: bool = False, upper: bool = False,
                 spacing: float = 0) -> int:
    measure = _measure(font, bold)
    limit = width * 0.97 / size          # in em, with a little slack for renderer differences
    extra = spacing / (size * PT_PER_PX)  # letter spacing, in em per character
    total = 0
    for p in paras:
        text = "".join(t for t, _ in p)
        text = text.upper() if upper else text
        if measure is None:
            per_em = 0.6 if font == MONO else (0.56 if bold else 0.5)
            total += max(1, math.ceil(len(text) * (per_em + extra) / limit))
            continue
        lines, cur = 1, 0.0
        space = measure(" ") + extra
        for word in text.split(" "):
            w = measure(word) + extra * len(word)
            if cur and cur + space + w > limit:
                lines, cur = lines + 1, w
            else:
                cur = cur + space + w if cur else w
        total += lines
    return total


def text_h(paras, size, width, lh=1.35, font=SANS, bold=False) -> float:
    return lines_needed(paras, size, width, font, bold) * size * lh


def font_px(node, default: float) -> float:
    m = re.search(r"font-size:\s*(\d+)px", node.get("style") or "")
    return float(m.group(1)) if m else default


# ---------- drawing ----------

class Canvas:
    def __init__(self, slide, theme):
        self.slide, self.theme = slide, theme

    def text(self, x, y, w, h, paras, size, color=None, bold=False, font=SANS, lh=1.35,
             align="left", anchor="top", upper=False, spacing=0, name=None):
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt
        if isinstance(paras, str):
            paras = [[(paras, {})]]
        box = self.slide.shapes.add_textbox(Emu(int(x * EMU_PER_PX)), Emu(int(y * EMU_PER_PX)),
                                            Emu(int(w * EMU_PER_PX)), Emu(int(h * EMU_PER_PX)))
        if name:
            box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
            p.line_spacing = Pt(size * PT_PER_PX * lh)
            for chunk, st in para:
                r = p.add_run()
                r.text = chunk.upper() if upper else chunk
                f = r.font
                f.name = st.get("font", font)
                f.size = Pt(round(size * PT_PER_PX * 2) / 2)
                f.bold = st.get("bold", bold)
                f.color.rgb = RGBColor.from_string(st.get("color", color or self.theme["ink"]))
                if spacing:
                    r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
        return box

    def box(self, x, y, w, h, fill=CARD, line=RULE, radius=10, line_w=1.0):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu, Pt
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        s = self.slide.shapes.add_shape(kind, Emu(int(x * EMU_PER_PX)), Emu(int(y * EMU_PER_PX)),
                                        Emu(int(w * EMU_PER_PX)), Emu(int(h * EMU_PER_PX)))
        if radius:
            s.adjustments[0] = min(0.5, radius / max(1.0, min(w, h)))
        s.shadow.inherit = False
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(line_w)
        else:
            s.line.fill.background()
        return s

    def rule(self, x, y, w, color=RULE, weight=1.0):
        return self.box(x, y, w, max(1.0, weight), fill=color, line=None, radius=0)


# ---------- components: each returns its height, and draws only when c is a Canvas ----------

def kicker(c, el, y, theme):
    if c:
        c.text(LEFT, y, CONTENT_W, 30, runs_of(el), 22, theme["accent"], font=MONO, lh=1.0, upper=True, spacing=2)
    return 24


def heading(c, el, y, theme):
    default, lh = (112, 0.98) if el.name == "h1" else (66, 1.04)
    size = font_px(el, default)
    paras = runs_of(el)
    h = text_h(paras, size, CONTENT_W, lh + 0.06, bold=True)
    if c:
        c.text(LEFT, y, CONTENT_W, h, paras, size, theme["ink"], bold=True, lh=lh)
    return h


def paragraph(c, el, y, theme):
    classes = el.get("class") or []
    lead = "lead" in classes
    size = font_px(el, 36 if lead else 28)
    font = MONO if "mono" in classes else SANS
    width = min(CONTENT_W, size * (34 if lead else 40)) if font == SANS else CONTENT_W
    paras = runs_of(el)
    lh = 1.32 if lead else 1.4
    h = text_h(paras, size, width, lh, font)
    if c:
        c.text(LEFT, y, width, h, paras, size, theme["ink"] if lead else theme["muted"], font=font, lh=lh)
    return h


def _stack(c, x, y, w, parts, draw=True):
    """parts: (paras, size, color, font, bold, lh, gap_after, upper). Returns total height."""
    cur = y
    for paras, size, color, font, bold, lh, gap, upper in parts:
        h = text_h(paras, size, w, lh, font, bold)
        if c and draw:
            c.text(x, cur, w, h, paras, size, color, bold=bold, font=font, lh=lh, upper=upper,
                   spacing=1.5 if upper else 0)
        cur += h + gap
    return cur - y - (parts[-1][6] if parts else 0)


def _door_parts(door):
    parts = []
    for child in door.find_all(["div", "p"], recursive=False):
        classes = child.get("class") or []
        if "who" in classes:
            parts.append((runs_of(child), 21, "0B6E8A", MONO, False, 1.0, 14, True))
        elif "m" in classes:
            parts.append((runs_of(child), 26, "566977", SANS, False, 1.35, 12, False))
        else:
            parts.append((runs_of(child), 26, INK, MONO if "mono" in classes else SANS, False, 1.35, 12, False))
    return parts


def layer(c, el, x, y, w):
    size = font_px(el, 26)
    span = el.find("span")
    label = [[(t, s) for t, s in p] for p in runs_of(el)] if not span else None
    if span:
        head = "".join(t for t in el.find_all(string=True, recursive=False)).strip()
        label = [[(head + "    ", {}), (span.get_text(" ", strip=True), {"color": "9DB0BD"})]]
    inner = w - 68
    h = text_h(label, size, inner, 1.45, MONO) + 48
    if c:
        c.box(x, y, w, h, fill=INK, line=None)
        c.text(x + 34, y + 24, inner, h - 48, label, size, PAPER, font=MONO, lh=1.45)
    return h


def cols(c, el, y, theme):
    n = int(re.search(r"c(\d)", " ".join(el.get("class"))).group(1))
    gap = 36
    doors = [d for d in el.find_all("div", recursive=False) if "door" in (d.get("class") or [])]
    w = (CONTENT_W - gap * (n - 1)) / n
    inner = w - 68
    h = max(_stack(None, 0, 0, inner, _door_parts(d)) for d in doors) + 60
    for i, d in enumerate(doors):
        x = LEFT + i * (w + gap)
        if c:
            c.box(x, y, w, h)
            _stack(c, x + 34, y + 30, inner, _door_parts(d))
    total = h
    for bar in [d for d in el.find_all("div", recursive=False) if "layer" in (d.get("class") or [])]:
        total += gap + layer(c, bar, LEFT, y + total + gap, CONTENT_W)
    return total


def email(c, el, y, theme):
    w, pad = 1240, 44
    inner = w - 2 * pad
    hdr = el.find(class_="hdr")
    hdr_paras = [[(t, dict(s, color=INK) if s.get("bold") else s) for t, s in p] for p in runs_of(hdr)]
    hdr_paras = [[(t.replace("\xa0", " "), dict(s, bold=False)) for t, s in p] for p in hdr_paras]
    hdr_h = text_h(hdr_paras, 21, inner, 1.5, MONO)
    body = [runs_of(p) for p in el.find_all("p", recursive=False)]
    body_h = sum(text_h(p, 29, inner, 1.42) for p in body) + 18 * (len(body) - 1)
    h = 36 + hdr_h + 16 + 18 + body_h + 36
    if c:
        c.box(LEFT, y, w, h)
        c.text(LEFT + pad, y + 36, inner, hdr_h, hdr_paras, 21, "566977", font=MONO, lh=1.5)
        c.rule(LEFT + pad, y + 36 + hdr_h + 16, inner)
        cur = y + 36 + hdr_h + 16 + 18
        for p in body:
            ph = text_h(p, 29, inner, 1.42)
            c.text(LEFT + pad, cur, inner, ph, p, 29, INK, lh=1.42)
            cur += ph + 18
    return h


def roles(c, el, y, theme):
    items = el.find_all("div", class_="role")
    gap = 16
    w = (CONTENT_W - gap * (len(items) - 1)) / len(items)
    def parts(r):
        return [(runs_of(r.find("b")), 27, INK, SANS, True, 1.2, 6, False),
                (runs_of(r.find("span")), 24, "566977", SANS, False, 1.3, 0, False)]
    h = max(_stack(None, 0, 0, w - 8, parts(r)) for r in items) + 18
    for i, r in enumerate(items):
        x = LEFT + i * (w + gap)
        if c:
            c.rule(x, y, w)
            _stack(c, x, y + 18, w - 8, parts(r))
    return h


def flow(c, el, y, theme):
    cells = el.find_all("div", recursive=False)
    w = CONTENT_W / len(cells)
    tint = {"ai": "0B6E8A", "py": "566977", "you": ALERT}
    body_h = max(text_h(runs_of(cell.find("span")), 22, w - 44, 1.35) for cell in cells)
    h = max(230, 76 + body_h + 20 + 48)
    if c:
        c.box(LEFT, y, CONTENT_W, h)
    for i, cell in enumerate(cells):
        x = LEFT + i * w
        if c:
            if i:
                c.rule(x, y, 1, RULE)
                c.slide.shapes[-1].height = int(h * EMU_PER_PX)
            c.text(x + 22, y + 26, w - 44, 40, runs_of(cell.find("b")), 32, INK, bold=True, lh=1.1)
            c.text(x + 22, y + 76, w - 44, body_h, runs_of(cell.find("span")), 22, "566977", lh=1.35)
            c.text(x + 22, y + h - 48, w - 44, 24, runs_of(cell.find("em")), 19,
                   tint[(cell.get("class") or ["py"])[0]], font=MONO, lh=1.0, upper=True, spacing=1.5)
    return h


def _cell_border(cell, color, weight_pt, fill=None):
    """python-pptx has no cell-border API; write the bottom line and clear the other three."""
    from lxml import etree
    from pptx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    made = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        if tag == "a:lnB":
            ln.set("w", str(int(weight_pt * 12700)))
            solid = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(solid, qn("a:srgbClr")).set("val", color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        made.append(ln)
    # Schema order inside tcPr: borders first, then the fill.
    for ln in reversed(made):
        tcPr.insert(0, ln)


def table(c, el, y, theme):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt
    small = "sm" in (el.get("class") or [])
    size, pad_v = (25, 15) if small else (30, 19)
    head = el.find("thead").find_all("th")
    rows = el.find("tbody").find_all("tr")
    ncol = len(head)
    widths = [0.22, 0.39, 0.39] if small else [0.46, 0.2, 0.34]
    col_w = [CONTENT_W * f for f in widths[:ncol]]
    row_h = [36.0]
    for tr in rows:
        cells = tr.find_all("td")
        row_h.append(max(text_h(runs_of(td), size, col_w[i] - 36, 1.25, bold=bool(td.find("b")))
                         for i, td in enumerate(cells)) + 2 * pad_v)
    total = sum(row_h)
    if not c:
        return total
    shape = c.slide.shapes.add_table(len(rows) + 1, ncol, Emu(int(LEFT * EMU_PER_PX)), Emu(int(y * EMU_PER_PX)),
                                     Emu(int(CONTENT_W * EMU_PER_PX)), Emu(int(total * EMU_PER_PX)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    style = shape._element.graphic.graphicData.tbl.tblPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Emu(int(cw * EMU_PER_PX))
    for r, rh in enumerate(row_h):
        tbl.rows[r].height = Emu(int(rh * EMU_PER_PX))

    def fill(cell, paras, fsize, color, font=SANS, bold=False, right=False, upper=False, bg=None):
        cell.margin_left = cell.margin_right = Emu(18 * EMU_PER_PX)
        cell.margin_top = cell.margin_bottom = Emu(int(pad_v * 0.6 * EMU_PER_PX))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if bg:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(bg)
        else:
            cell.fill.background()
        tf = cell.text_frame
        tf.word_wrap = True
        for k, para in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
            for chunk, st in para:
                run = p.add_run()
                run.text = chunk.upper() if upper else chunk
                run.font.name = st.get("font", font)
                run.font.size = Pt(round(fsize * PT_PER_PX * 2) / 2)
                run.font.bold = st.get("bold", bold)
                run.font.color.rgb = RGBColor.from_string(st.get("color", color))

    for i, th in enumerate(head):
        cell = tbl.cell(0, i)
        fill(cell, runs_of(th), 20, "566977", font=MONO, right="n" in (th.get("class") or []), upper=True)
        _cell_border(cell, INK, 1.5)
    for r, tr in enumerate(rows, start=1):
        gap_row = "gap" in (tr.get("class") or [])
        for i, td in enumerate(tr.find_all("td")):
            right = "n" in (td.get("class") or [])
            hot = gap_row and i == 1
            cell = tbl.cell(r, i)
            fill(cell, runs_of(td), size, ALERT if hot else INK, bold=hot, right=right,
                 bg=ALERT_SOFT if gap_row else None)
            _cell_border(cell, RULE, 0.75)
    return total


def props(c, el, y, theme):
    cards = el.find_all("div", class_="prop", recursive=False)
    small = [p for p in cards if "agent" not in p.get("class")]
    gap = 16
    w = (CONTENT_W - gap * (len(small) - 1)) / len(small)
    def parts(p):
        return [(runs_of(p.find(class_="c")), 58, "0B6E8A", SANS, True, 1.0, 8, False),
                (runs_of(p.find(class_="t")), 21, INK, MONO, False, 1.25, 8, False),
                (runs_of(p.find(class_="d")), 21, "566977", SANS, False, 1.3, 0, False)]
    h = max(_stack(None, 0, 0, w - 48, parts(p)) for p in small) + 44
    for i, p in enumerate(small):
        x = LEFT + i * (w + gap)
        if c:
            c.box(x, y, w, h)
            _stack(c, x + 24, y + 22, w - 48, parts(p))
    total = h
    for p in [p for p in cards if "agent" in p.get("class")]:
        title, desc = runs_of(p.find(class_="t")), runs_of(p.find(class_="d"))
        dw = CONTENT_W - 48 - 290 - 30
        ah = max(60, text_h(desc, 24, dw, 1.3)) + 44
        if c:
            c.box(LEFT, y + total + gap, CONTENT_W, ah, fill=INK, line=None)
            c.text(LEFT + 24, y + total + gap, 290, ah, title, 30, PAPER, font=MONO, lh=1.1, anchor="middle")
            c.text(LEFT + 24 + 290 + 30, y + total + gap, dw, ah, desc, 24, "BFCBD3", lh=1.3, anchor="middle")
        total += gap + ah
    return total


def choices(c, el, y, theme):
    cards = el.find_all("div", class_="choice")
    gap = 20
    w = (CONTENT_W - gap) / 2
    def parts(ch):
        return [(runs_of(ch.find("b")), 36, INK, SANS, True, 1.15, 14, False),
                (runs_of(ch.find("p")), 28, "566977", SANS, False, 1.35, 14, False),
                (runs_of(ch.find(class_="def")), 22, "0B6E8A", MONO, False, 1.3, 0, False)]
    total = 0.0
    for r in range(0, len(cards), 2):
        pair = cards[r:r + 2]
        h = max(_stack(None, 0, 0, w - 68, parts(ch)) for ch in pair) + 60
        for i, ch in enumerate(pair):
            x = LEFT + i * (w + gap)
            if c:
                c.box(x, y + total, w, h)
                _stack(c, x + 34, y + total + 30, w - 68, parts(ch))
        total += h + gap
    return total - gap


def sms(c, el, y, theme):
    m = re.search(r"max-width:\s*(\d+)px", el.get("style") or "")
    w = float(m.group(1)) if m else CONTENT_W
    hit = "hit" in (el.get("class") or [])
    parts = [(runs_of(el.find("time")), 22, "566977", MONO, False, 1.0, 10, False),
             (runs_of(el.find("p")), 34, INK, SANS, False, 1.3, 0, False)]
    h = _stack(None, 0, 0, w - 36, parts) + 28
    if c:
        c.box(LEFT, y, w, h, fill=ALERT_SOFT if hit else CARD, line=ALERT if hit else RULE, radius=16,
              line_w=1.5 if hit else 1.0)
        _stack(c, LEFT + 18, y + 14, w - 36, parts)
    return h


def versus(c, el, y, theme):
    cards = el.find_all("div", class_="vs")
    gap = 24
    w = (CONTENT_W - gap) / 2
    def parts(v, bad):
        return [(runs_of(v.find(class_="thr")), 24, "566977", MONO, False, 1.0, 10, False),
                (runs_of(v.find(class_="res")), 54, ALERT if bad else INK, SANS, True, 1.05, 8, False),
                (runs_of(v.find("p")), 27, "566977", SANS, False, 1.35, 0, False)]
    h = max(_stack(None, 0, 0, w - 64, parts(v, False)) for v in cards) + 56
    for i, v in enumerate(cards):
        bad = "bad" in v.get("class")
        x = LEFT + i * (w + gap)
        if c:
            c.box(x, y, w, h, fill=ALERT_SOFT if bad else CARD, line=ALERT if bad else RULE,
                  line_w=1.5 if bad else 1.0)
            _stack(c, x + 32, y + 28, w - 64, parts(v, bad))
    return h


def push(c, el, y, theme):
    total = 0.0
    width = min(CONTENT_W, 35 * 38) - 64
    for n, li in enumerate(el.find_all("li"), start=1):
        hot = "hot" in (li.get("class") or [])
        paras = runs_of(li)
        h = text_h(paras, 35, width, 1.3)
        if c:
            c.text(LEFT, y + total + 6, 56, 36, str(n), 26, ALERT if hot else theme["accent"], font=MONO, lh=1.0)
            c.text(LEFT + 64, y + total, width, h, paras, 35, ALERT if hot else theme["ink"], lh=1.3)
        total += h + 24
    return total - 24


def themes(c, el, y, theme):
    total, width = 0.0, 1240
    for li in el.find_all("li"):
        parts = [(runs_of(li.find("b")), 40, theme["ink"], SANS, True, 1.12, 6, False),
                 (runs_of(li.find("span")), 25, theme["muted"], SANS, False, 1.35, 0, False)]
        h = _stack(c, LEFT, y + total, width, parts)
        total += h + 22
    return total - 22


def component(el):
    classes = el.get("class") or []
    if el.name in ("h1", "h2"):
        return heading
    if el.name == "p":
        return paragraph
    if el.name == "table":
        return table
    if el.name == "ul":
        return push if "push" in classes else themes
    for key, fn in (("kicker", kicker), ("cols", cols), ("email", email), ("roles", roles), ("flow", flow),
                    ("props", props), ("choices", choices), ("sms", sms), ("versus", versus)):
        if key in classes:
            return fn
    if "layer" in classes:
        return lambda c, e, y, theme: layer(c, e, LEFT, y, CONTENT_W)
    raise ValueError(f"no editable renderer for <{el.name} class={classes}>")


# ---------- the footer rail and the slide loop ----------

def segments(source: str) -> list[tuple[str, int, int]]:
    raw = re.search(r"var SEGS=(\[\[.*?\]\]);", source).group(1)
    return [(n, int(a), int(b)) for n, a, b in re.findall(r'\["([^"]+)",(\d+),(\d+)\]', raw)]


def rail(c, theme, segs, seg, pre, index, count):
    name, a, b = ("Before we start", 0, 0) if pre else segs[seg]
    label = name if pre else f"{a}–{b} min · {name}"
    y = H - 44 - 19
    c.text(LEFT, y, 330, 22, label, 19, theme["muted"], font=MONO, lh=1.0, name="Timeline label")
    c.text(W - RIGHT - 110, y, 110, 22, f"{index} / {count}", 19, theme["muted"], font=MONO, lh=1.0, align="right",
           name="Slide number")
    x0, x1, gap = LEFT + 350, W - RIGHT - 130, 5
    span = sum(e - s for _, s, e in segs)
    usable = (x1 - x0) - gap * (len(segs) - 1)
    x = x0
    for k, (_, s, e) in enumerate(segs):
        w = usable * (e - s) / span
        active = (not pre) and k == seg
        past = (not pre) and k < seg
        c.box(x, y + 6, w, 8, fill=theme["accent"] if active else (theme["muted"] if past else theme["rail"]),
              line=None, radius=0)
        x += w + gap


def build(source: str, out: Path) -> Path:
    from bs4 import BeautifulSoup
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu

    soup = BeautifulSoup(source, "html.parser")
    sections = soup.find("main", id="deck").find_all("section", class_="slide", recursive=False)
    meta, segs = slides_meta(source), segments(source)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * EMU_PER_PX), Emu(H * EMU_PER_PX)
    for index, (section, m) in enumerate(zip(sections, meta), start=1):
        classes = section.get("class")
        theme = DARK if "invert" in classes else LIGHT
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme["bg"])
        c = Canvas(slide, theme)
        parts = [e for e in section.find_all(True, recursive=False) if e.name != "aside"]
        split = next((i for i, e in enumerate(parts) if "spacer" in (e.get("class") or [])), None)
        upper, lower = (parts, []) if split is None else (parts[:split], parts[split + 1:])
        lower = [e for e in lower if "hint" not in (e.get("class") or [])]   # browser keyboard hints

        def height(elements):
            return sum(component(e)(None, e, 0, theme) for e in elements) + GAP * max(0, len(elements) - 1)

        y = TOP
        if "center" in classes and not lower:
            y = TOP + max(0.0, (H - TOP - BOTTOM - height(upper)) / 2)
        for e in upper:
            y += component(e)(c, e, y, theme) + GAP
        y = H - BOTTOM - height(lower)
        for e in lower:
            y += component(e)(c, e, y, theme) + GAP
        rail(c, theme, segs, int(section.get("data-seg", 0)), section.has_attr("data-pre"), index, len(sections))
        slide.notes_slide.notes_text_frame.text = m["notes"]
    prs.save(str(out))
    return out


def main() -> Path:
    out = build(SOURCE.read_text(encoding="utf-8"), OUT)
    print(f"wrote {out}")
    return out


if __name__ == "__main__":
    main()
