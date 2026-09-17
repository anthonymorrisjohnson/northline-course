"""Export slides/present.html to slides/present.pptx: one full-bleed image per slide, speaker notes attached.

Author-only. Needs Google Chrome (or Edge or Chromium) and python-pptx:

    uv run --with python-pptx python scripts/export_deck.py

The slides are images, so they look exactly like the browser deck in PowerPoint, Keynote, and
Google Slides. To change the words, edit slides/present.html and run this again.
"""
import html
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "slides" / "present.html"
OUT = ROOT / "slides" / "present.pptx"
SHELL = ('<!doctype html><html lang="en"><head><meta charset="utf-8">'
         '<meta name="viewport" content="width=device-width, initial-scale=1">'
         '<style>html,body{margin:0}</style></head><body>%s</body></html>')
BROWSERS = ["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
            "/Applications/Chromium.app/Contents/MacOS/Chromium",
            "google-chrome", "chromium", "chromium-browser", "msedge"]


def _text(fragment: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", fragment))).strip()


def slides_meta(source: str) -> list[dict]:
    """Title and plain-text speaker notes for every slide, in order."""
    out = []
    for attrs, body in re.findall(r'<section class="slide[^"]*"([^>]*)>(.*?)</section>', source, re.S):
        title = re.search(r'data-title="([^"]*)"', attrs)
        notes = re.search(r'<aside class="notes">(.*?)</aside>', body, re.S)
        lines = []
        for label, rest in re.findall(r"<p><b>(.*?)</b>(.*?)</p>", notes.group(1) if notes else "", re.S):
            lines.append(f"{_text(label).upper()}: {_text(rest)}")
        out.append({"title": html.unescape(title.group(1)) if title else "", "notes": "\n\n".join(lines)})
    return out


def find_browser() -> str:
    for b in BROWSERS:
        if Path(b).exists() or shutil.which(b):
            return b if Path(b).exists() else shutil.which(b)
    raise SystemExit("No Chrome, Edge, or Chromium found. Install one, or present from /slides in the browser.")


def render(count: int, page: Path, out_dir: Path, browser: str) -> list[Path]:
    shots = []
    for n in range(1, count + 1):
        png = out_dir / f"slide-{n:02d}.png"
        subprocess.run([browser, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                        "--window-size=1600,900", "--virtual-time-budget=6000", f"--screenshot={png}",
                        f"{page.as_uri()}?export=1#/{n}"], check=True, capture_output=True, timeout=120)
        shots.append(png)
        print(f"rendered {n}/{count}")
    return shots


def build(shots: list[Path], meta: list[dict], out: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for png, m in zip(shots, meta):
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic._element.xpath(".//p:cNvPr")[0].set("descr", m["title"])   # alt text for screen readers
        slide.notes_slide.notes_text_frame.text = m["notes"]
    prs.save(str(out))
    return out


def main(keep: Path | None = None) -> Path:
    source = SOURCE.read_text(encoding="utf-8")
    meta = slides_meta(source)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        page = tmp / "deck.html"
        page.write_text(SHELL % source, encoding="utf-8")
        shots = render(len(meta), page, tmp, find_browser())
        build(shots, meta, OUT)
        if keep:
            keep.mkdir(parents=True, exist_ok=True)
            for s in shots:
                shutil.copy(s, keep / s.name)
    print(f"wrote {OUT} ({len(meta)} slides)")
    return OUT


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else None)
