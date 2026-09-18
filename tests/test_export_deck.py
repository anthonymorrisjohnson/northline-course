from scripts import export_deck as ex


def test_slides_meta_reads_titles_and_notes():
    meta = ex.slides_meta(ex.SOURCE.read_text(encoding="utf-8"))
    assert len(meta) == 16
    assert meta[0]["title"] == "Setup, while you arrive" and meta[2]["title"] == "The Monday email"
    assert all(m["notes"] for m in meta)
    assert meta[2]["notes"].startswith("ROOM: Read the email aloud")
    assert "<" not in meta[5]["notes"] and "\n\n" in meta[5]["notes"]


def test_slides_meta_handles_a_slide_without_notes():
    meta = ex.slides_meta('<section class="slide" data-title="A &amp; B"><h2>x</h2></section>')
    assert meta == [{"title": "A & B", "notes": ""}]


def test_follow_along_is_built_from_its_body_and_has_notes():
    from scripts import build_follow_along as b
    built = b.OUT.read_text(encoding="utf-8")
    assert built.startswith("<title>Northline Follow Along</title>")
    b.build()
    assert b.OUT.read_text(encoding="utf-8") == built, "follow-along.html is stale: run scripts/build_follow_along.py"
    meta = ex.slides_meta(built)
    assert len(meta) == 21 and all(m["notes"] for m in meta) and meta[10]["title"] == "Step 1: /pm-run"
    assert meta[7]["title"].startswith("Discuss:")
    assert [m["title"] for m in meta[3:7]] == [t for t in (m["title"] for m in meta) if t.startswith("Basics:")]


def test_handout_carries_every_exhibit_and_matches_the_triage_data():
    import json
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    handout = (root / "docs" / "exhibits.md").read_text(encoding="utf-8")
    for letter in "ABCDE":
        assert f"## Exhibit {letter}:" in handout
    for row in json.loads((root / "northline/agents/triage/exhibit_e.json").read_text(encoding="utf-8")):
        assert f"| {row['n']} | {row['time']} | {row['text']} |" in handout
    # The pocket memo is held back until minute 12, so it must print on its own, last.
    assert handout.rindex('<div class="pagebreak"></div>') < handout.index("## Exhibit D:")
    assert handout.index("## Exhibit D:") > handout.index("## Exhibit E:")
    assert "answer key" not in handout.lower() and "Urgent clinical" not in handout
    assert (root / "docs" / "exhibits.pdf").stat().st_size > 10_000


def test_editable_deck_is_native_text_tables_and_notes(tmp_path):
    import pytest
    pytest.importorskip("pptx")
    pytest.importorskip("bs4")
    from pathlib import Path
    from pptx import Presentation
    from scripts import export_deck_editable as ed

    source = (Path(__file__).resolve().parents[1] / "slides" / "present.html").read_text(encoding="utf-8")
    prs = Presentation(str(ed.build(source, tmp_path / "deck.pptx")))
    meta = ex.slides_meta(source)
    assert len(prs.slides) == len(meta)
    for slide, m in zip(prs.slides, meta):
        assert not [s for s in slide.shapes if s.shape_type == 13], "pictures are not editable"
        assert slide.notes_slide.notes_text_frame.text == m["notes"]
    words = " ".join(s.text_frame.text for s in prs.slides[4].shapes if s.has_text_frame)
    assert "Two front doors, one set of tools" in words and "2,400 flags a week" in words
    assert sum(1 for slide in prs.slides for s in slide.shapes if s.has_table) == 2
